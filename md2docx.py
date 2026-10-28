#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 Markdown 报告转换为 Word（.docx）。

用法：
    <python> md2docx.py <input.md> <output.docx>

支持：标题层级 / 段落 / 引用块 / GFM 表格（表头底纹+跨页重复）/ 代码块（灰底等宽）/
      有序与无序列表 / 任务清单 / 粗体 / 行内代码 / 链接 / 分隔线 / 页脚页码 / A4 页边距。
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "微软雅黑"
BODY_FONT = "等线"
MONO_FONT = "Consolas"

INK = "1F2937"
INK_SOFT = "475569"
INK_FAINT = "94A3B8"
HEAD_COLOR = {1: "0B3B33", 2: "0F766E", 3: "334155"}
LINK_COLOR = "1155CC"
CODE_INK = "9A3412"
QUOTE_BG = "F8FAFC"
CODE_BG = "F1F5F9"
TABLE_HEAD_BG = "E8F5F1"

AVAILABLE_CM = 16.6  # A4(21cm) - 左右边距(2.2cm*2)


# ----------------------------------------------------------------------------
# 低层工具
# ----------------------------------------------------------------------------
def set_run(run, size=10.5, bold=False, italic=False, color=INK, mono=False,
            underline=False, ea_font=None):
    """统一设置 run 的字体（含中文 eastAsia 字体，缺了中文会掉字形）。"""
    name = MONO_FONT if mono else BODY_FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.name = name
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    rf.set(qn("w:eastAsia"), ea_font or (name if not mono else CN_FONT))
    return run


def style_fonts(style, ascii_name, ea_name):
    rPr = style.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), ascii_name)
    rf.set(qn("w:hAnsi"), ascii_name)
    rf.set(qn("w:eastAsia"), ea_name)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def shade_paragraph(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def left_bar(par, color="99F6E4", size="18"):
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    pbdr.append(left)
    pPr.append(pbdr)


def bottom_rule(par, color="CBD5E1"):
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def add_field(par, instr):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    par._p.append(fld)


# ----------------------------------------------------------------------------
# 行内格式：**粗体** / `代码` / [文字](链接)
# ----------------------------------------------------------------------------
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_hyperlink(par, target, text, size=10.5):
    """插入真正的超链接（Word 里可点击）。相对路径会按文档所在目录解析。"""
    try:
        r_id = par.part.relate_to(target, RT.HYPERLINK, is_external=True)
    except Exception:
        set_run(par.add_run(text), size=size, color=LINK_COLOR, underline=True)
        return
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), BODY_FONT)
    rf.set(qn("w:hAnsi"), BODY_FONT)
    rf.set(qn("w:eastAsia"), BODY_FONT)
    rPr.append(rf)
    for tag, val in (("w:color", LINK_COLOR), ("w:sz", str(int(size * 2))), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rPr.append(el)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    hl.append(r)
    par._p.append(hl)


def add_inline(par, text, size=10.5, bold=False, color=INK):
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_inline(par, tok[2:-2], size=size, bold=True, color=color)   # 递归处理粗体里嵌的链接/代码
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            set_run(par.add_run(tok[1:-1]), size=size - 0.5, color=CODE_INK, mono=True)
        elif tok.startswith("[") and "](" in tok:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            label, target = m.group(1), m.group(2)
            if target.startswith("#"):
                set_run(par.add_run(label), size=size, color=LINK_COLOR, underline=True)
            else:
                add_hyperlink(par, target, label, size=size)
        else:
            set_run(par.add_run(tok), size=size, bold=bold, color=color)


# ----------------------------------------------------------------------------
# Markdown 解析
# ----------------------------------------------------------------------------
SEP_RE = re.compile(r"^:?-{2,}:?$")


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse_md(text):
    blocks, lines, i = [], text.split("\n"), 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.lstrip().startswith("```"):
            code, i = [], i + 1
            while i < len(lines) and not lines[i].lstrip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1
            blocks.append(("code", code))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", line.strip()):
            blocks.append(("hr",))
            i += 1
            continue
        if line.lstrip().startswith("|"):
            raw = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw.append(split_row(lines[i]))
                i += 1
            rows = [r for r in raw if not all(SEP_RE.match(c) for c in r if c)]
            if rows:
                blocks.append(("table", rows))
            continue
        if line.lstrip().startswith(">"):
            quote = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote.append(lines[i].lstrip()[1:].strip())
                i += 1
            blocks.append(("quote", quote))
            continue
        m = re.match(r"^\s*[-*+]\s+\[( |x|X)\]\s*(.*)$", line)
        if m:
            blocks.append(("task", m.group(2).strip()))
            i += 1
            continue
        m = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if m:
            blocks.append(("ul", m.group(1).strip()))
            i += 1
            continue
        m = re.match(r"^\s*(\d+)[.、)]\s+(.*)$", line)
        if m:
            blocks.append(("ol", m.group(1), m.group(2).strip()))
            i += 1
            continue
        blocks.append(("p", line.strip()))
        i += 1
    return blocks


# ----------------------------------------------------------------------------
# 文档骨架
# ----------------------------------------------------------------------------
def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    style_fonts(normal, BODY_FONT, BODY_FONT)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(4)

    title = doc.styles["Title"]
    title.font.size = Pt(19)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(HEAD_COLOR[1])
    style_fonts(title, CN_FONT, CN_FONT)
    title.paragraph_format.space_after = Pt(10)

    for lvl, (size, color) in {1: (14.5, HEAD_COLOR[2]), 2: (12, HEAD_COLOR[3]),
                               3: (11, HEAD_COLOR[3]), 4: (10.5, HEAD_COLOR[3])}.items():
        st = doc.styles["Heading %d" % lvl]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(color)
        style_fonts(st, CN_FONT, CN_FONT)
        st.paragraph_format.space_before = Pt(11)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("寝室公共开支智能记账与分摊提醒 · 第 "), size=8, color=INK_FAINT)
    add_field(footer, "PAGE")
    set_run(footer.add_run(" 页 / 共 "), size=8, color=INK_FAINT)
    add_field(footer, "NUMPAGES")
    set_run(footer.add_run(" 页"), size=8, color=INK_FAINT)
    return doc


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    weights = []
    for c in range(ncols):
        longest = max(len(re.sub(r"[*`]", "", rows[r][c])) for r in range(min(len(rows), 14)))
        weights.append(max(5, min(30, longest)))          # 限幅，避免某列吃掉整行
    total = float(sum(weights))
    widths = [Cm(AVAILABLE_CM * w / total) for w in weights]

    for r, row in enumerate(rows):
        for c, txt in enumerate(row):
            cell = table.cell(r, c)
            cell.width = widths[c]
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.line_spacing = 1.2
            add_inline(par, txt, size=8.5 if ncols > 3 else 9, bold=(r == 0))
            if r == 0:
                shade_cell(cell, TABLE_HEAD_BG)
    if len(rows) > 1:
        repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build(doc, blocks):
    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            _, level, txt = blk
            if level == 1:
                par = doc.add_paragraph(style="Title")
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(par, txt, size=19, bold=True, color=HEAD_COLOR[1])
            else:
                style = "Heading %d" % min(4, level - 1)
                par = doc.add_paragraph(style=style)
                add_inline(par, txt, size={2: 14.5, 3: 12, 4: 11}.get(level, 10.5),
                           bold=True, color=HEAD_COLOR.get(min(2, level - 1), INK))
        elif kind == "p":
            par = doc.add_paragraph()
            add_inline(par, blk[1])
        elif kind == "quote":
            for line in blk[1]:
                if not line:
                    continue
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.35)
                par.paragraph_format.space_after = Pt(3)
                par.paragraph_format.line_spacing = 1.35
                left_bar(par)
                shade_paragraph(par, QUOTE_BG)
                add_inline(par, line, size=9.5, color=INK_SOFT)
        elif kind == "code":
            for line in blk[1]:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.left_indent = Cm(0.3)
                par.paragraph_format.line_spacing = 1.05
                shade_paragraph(par, CODE_BG)
                set_run(par.add_run(line if line.strip() else " "),
                        size=8.5, mono=True, color="1F2937")
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif kind == "table":
            add_table(doc, blk[1])
        elif kind == "ul":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.75)
            par.paragraph_format.first_line_indent = Cm(-0.35)
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, "• " + blk[1])
        elif kind == "ol":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            par.paragraph_format.first_line_indent = Cm(-0.45)
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, blk[1] + ". " + blk[2])
        elif kind == "task":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            par.paragraph_format.first_line_indent = Cm(-0.45)
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, "☐ " + blk[1])
        elif kind == "hr":
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(7)
            bottom_rule(par)


def main():
    if len(sys.argv) < 3:
        print("用法: md2docx.py <input.md> <output.docx>")
        return 2
    src, dst = sys.argv[1], sys.argv[2]
    with open(src, "r", encoding="utf-8") as fh:
        md = fh.read()

    blocks = parse_md(md)
    doc = setup_document()
    build(doc, blocks)
    doc.save(dst)

    counts = {}
    for b in blocks:
        counts[b[0]] = counts.get(b[0], 0) + 1
    print("已生成: %s" % dst)
    print("块统计: %s" % ", ".join("%s=%d" % (k, v) for k, v in sorted(counts.items())))
    return 0


if __name__ == "__main__":
    sys.exit(main())
